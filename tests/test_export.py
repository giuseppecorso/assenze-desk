"""Export in Excel di giustificativi e pratiche (issue #3)."""

from datetime import date, datetime

from openpyxl import load_workbook

from assenze_desk.core.stato import Stato
from assenze_desk.storage.archivio import Pratica, RigaSalvata
from assenze_desk.storage.export import esporta_giustificativi, esporta_pratiche

CF = "FNTPRV80A01Z999Y"


def leggi(percorso):
    libro = load_workbook(percorso)
    foglio = libro.active
    righe = [[c.value for c in riga] for riga in foglio.iter_rows()]
    return foglio, righe


def riga(giorno, minuti):
    return RigaSalvata(CF, "Fittizio Luca", "201111111", "Alfa Srl", giorno, "ROL", "ROL", minuti)


def test_export_giustificativi(tmp_path):
    percorso = tmp_path / "export" / "giustificativi.xlsx"  # cartella creata se manca
    numero = esporta_giustificativi(
        [riga(date(2026, 10, 12), 90), riga(date(2026, 10, 13), None)], percorso
    )
    assert numero == 2

    foglio, righe = leggi(percorso)
    assert foglio.title == "Giustificativi"
    assert righe[0] == [
        "Giorno",
        "Codice fiscale",
        "Dipendente",
        "JobOrder",
        "Cliente",
        "Tipo",
        "Codice",
        "Ore",
    ]
    assert righe[1] == [
        datetime(2026, 10, 12),
        CF,
        "Fittizio Luca",
        "201111111",
        "Alfa Srl",
        "ROL",
        "ROL",
        1.5,
    ]
    assert righe[2][-1] is None  # ore da compilare: cella vuota
    assert foglio["A2"].number_format == "DD/MM/YYYY"
    assert foglio["H2"].number_format == "0.00"
    assert foglio["A1"].font.bold
    assert foglio.freeze_panes == "A2"
    assert foglio.auto_filter.ref == "A1:H3"


def test_export_vuoto_solo_intestazioni(tmp_path):
    percorso = tmp_path / "vuoto.xlsx"
    assert esporta_giustificativi([], percorso) == 0
    _foglio, righe = leggi(percorso)
    assert len(righe) == 1


def test_export_pratiche(tmp_path):
    pratiche = [
        Pratica(
            protocollo="411111111",
            cf=CF,
            dipendente="Fittizio Luca",
            joborder="201111111",
            cliente="Alfa Srl",
            inizio=date(2026, 10, 1),
            fine=date(2026, 10, 5),
            giorni_episodio=8,
            mail_lavoratore=True,
            note="continuazione",
        ),
        Pratica(protocollo="422222222"),
    ]
    percorso = tmp_path / "pratiche.xlsx"
    assert esporta_pratiche(pratiche, percorso) == 2

    foglio, righe = leggi(percorso)
    assert foglio.title == "Pratiche"
    intestazioni = righe[0]
    prima = dict(zip(intestazioni, righe[1], strict=True))
    assert prima["Protocollo"] == "411111111"
    assert (prima["Dal"], prima["Al"]) == (datetime(2026, 10, 1), datetime(2026, 10, 5))
    assert (prima["Giorni"], prima["Giorni episodio"]) == (5, 8)
    assert prima["Stato"] == Stato.PARZIALE.value  # 8 giorni di episodio, mail ufficio da fare
    assert (prima["Mail lavoratore"], prima["Mail ufficio"]) == ("Sì", "No")
    assert prima["Note"] == "continuazione"
    seconda = dict(zip(intestazioni, righe[2], strict=True))
    assert (seconda["Dal"], seconda["Giorni"], seconda["Stato"]) == (None, None, "Incompleta")
