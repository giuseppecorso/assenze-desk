"""Export in Excel: righe dei giustificativi (per il gestionale) e archivio delle pratiche.

Date come vere date di Excel, ore come numeri con due decimali, intestazioni in
grassetto, filtri e prima riga bloccata. Una cella vuota nelle ore vuol dire
"da compilare". Il file contiene dati personali: si salva fuori dal repo
(`.gitignore` esclude gli .xlsx).
"""

from collections.abc import Sequence
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from assenze_desk.storage.archivio import Pratica, RigaSalvata

_GIUSTIFICATIVI = [
    "Giorno",
    "Codice fiscale",
    "Dipendente",
    "JobOrder",
    "Cliente",
    "Tipo",
    "Codice",
    "Ore",
]
_PRATICHE = [
    "Protocollo",
    "Codice fiscale",
    "Dipendente",
    "JobOrder",
    "Cliente",
    "Dal",
    "Al",
    "Giorni",
    "Giorni episodio",
    "Stato",
    "Mail lavoratore",
    "Mail ufficio",
    "Note",
]


def esporta_giustificativi(righe: Sequence[RigaSalvata], percorso: Path) -> int:
    valori = [
        [
            r.giorno,
            r.cf,
            r.dipendente,
            r.joborder,
            r.cliente,
            r.tipo,
            r.codice,
            r.minuti / 60 if r.minuti is not None else None,
        ]
        for r in righe
    ]
    _salva(percorso, "Giustificativi", _GIUSTIFICATIVI, valori, date=["A"], numeri=["H"])
    return len(valori)


def esporta_pratiche(pratiche: Sequence[Pratica], percorso: Path) -> int:
    valori = [
        [
            p.protocollo,
            p.cf,
            p.dipendente,
            p.joborder,
            p.cliente,
            p.inizio,
            p.fine,
            (p.fine - p.inizio).days + 1 if p.inizio and p.fine else None,
            p.giorni_episodio,
            p.stato.value,
            "Sì" if p.mail_lavoratore else "No",
            "Sì" if p.mail_ufficio else "No",
            p.note,
        ]
        for p in pratiche
    ]
    _salva(percorso, "Pratiche", _PRATICHE, valori, date=["F", "G"], numeri=[])
    return len(valori)


def _salva(
    percorso: Path,
    titolo: str,
    intestazioni: list[str],
    righe: list[list[Any]],
    date: list[str],
    numeri: list[str],
) -> None:
    libro = Workbook()
    foglio = libro.active
    foglio.title = titolo
    foglio.append(intestazioni)
    for riga in righe:
        foglio.append(riga)

    for cella in foglio[1]:
        cella.font = Font(bold=True)
    for colonna in date:
        for cella in foglio[colonna][1:]:
            cella.number_format = "DD/MM/YYYY"
    for colonna in numeri:
        for cella in foglio[colonna][1:]:
            cella.number_format = "0.00"
    ultima = get_column_letter(len(intestazioni))
    foglio.auto_filter.ref = f"A1:{ultima}{len(righe) + 1}"
    foglio.freeze_panes = "A2"
    for numero, intestazione in enumerate(intestazioni, start=1):
        foglio.column_dimensions[get_column_letter(numero)].width = max(12, len(intestazione) + 4)

    Path(percorso).parent.mkdir(parents=True, exist_ok=True)
    libro.save(percorso)
