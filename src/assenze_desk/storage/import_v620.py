"""Import dei dati della v6.20: i due file Excel che la vecchia versione aggiornava.

- Tracking delle malattie: una riga per pratica (colonne Protocollo, CF,
  DataInizio, DataFine, MailLav, MailUff, StatoManuale, Stato, ...). Lo stato
  si importa solo se era stato messo a mano: quello automatico lo ricalcolano
  le regole nuove. Serve anche all'episodio: i certificati vecchi dello stesso
  CF contano per le continuazioni.
- Giustificativi: una riga per giorno (CodiceFiscale, Giorno,
  TipoGiustificativo, QuantitaOre in ore). La v6.20 non segnava le righe
  esportate, quindi le vecchie si importano come già esportate: non tornano
  nell'export nuovo. Una riga con più giorni non si divide a indovinare.

Il tipo di file si riconosce dalle colonne, non dal nome. Le pratiche e le
righe già presenti nell'archivio non si toccano: importare due volte non crea
doppioni. Valori illeggibili = campo vuoto, mai inventato.
"""

from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Literal

from openpyxl import load_workbook

from assenze_desk.core.date import leggi_data
from assenze_desk.core.stato import Stato
from assenze_desk.core.testo import normalizza
from assenze_desk.storage.archivio import Archivio, Pratica, RigaSalvata

_SI = {"si", "s", "yes", "y", "1", "true", "vero"}


@dataclass(frozen=True)
class DatiV620:
    tipo: Literal["malattie", "giustificativi"]
    pratiche: tuple[Pratica, ...] = ()
    righe: tuple[RigaSalvata, ...] = ()
    avvisi: tuple[str, ...] = ()


@dataclass(frozen=True)
class EsitoImport:
    nuove: int
    gia_presenti: int


def leggi_file_v620(percorso: Path) -> DatiV620:
    libro = load_workbook(percorso, read_only=True, data_only=True)
    try:
        for foglio in libro.worksheets:
            righe = foglio.iter_rows(values_only=True)
            colonne = {_chiave(h): i for i, h in enumerate(next(righe, ())) if h is not None}
            if {"protocollo", "datainizio", "cf"} <= colonne.keys():
                return _malattie(colonne, righe)
            if {"codicefiscale", "giorno", "tipogiustificativo"} <= colonne.keys():
                return _giustificativi(colonne, righe)
    finally:
        libro.close()
    raise ValueError(
        "il file non sembra un file della versione precedente "
        "(tracking delle malattie o giustificativi)"
    )


def importa_in_archivio(archivio: Archivio, dati: DatiV620) -> EsitoImport:
    nuove = gia_presenti = 0
    for pratica in dati.pratiche:
        if archivio.pratica(pratica.protocollo) is not None:
            gia_presenti += 1
            continue
        archivio.salva_pratica(pratica)
        nuove += 1

    presenti = {(r.cf, r.giorno, r.tipo) for r in archivio.righe()}
    da_aggiungere = []
    for riga in dati.righe:
        chiave = (riga.cf, riga.giorno, riga.tipo)
        if chiave in presenti:
            gia_presenti += 1
            continue
        presenti.add(chiave)
        da_aggiungere.append(riga)
    archivio.aggiungi_righe(da_aggiungere)
    nuove += len(da_aggiungere)

    cosa = "pratiche" if dati.tipo == "malattie" else "righe dei giustificativi"
    archivio.registra(
        f"Importate {cosa} dalla versione precedente",
        dettaglio=f"{nuove} nuove, {gia_presenti} già presenti",
    )
    return EsitoImport(nuove, gia_presenti)


# --- Lettura ---


def _malattie(colonne: dict[str, int], righe: Any) -> DatiV620:
    pratiche: dict[str, Pratica] = {}
    avvisi = []
    for numero, riga in enumerate(righe, start=2):
        valore = _lettore(colonne, riga)
        if not any(v not in (None, "") for v in riga):
            continue
        protocollo = _testo(valore("protocollo"))
        if not protocollo:
            avvisi.append(f"riga {numero}: manca il protocollo, non importata")
            continue
        stato_manuale = None
        if _si(valore("statomanuale")):
            stato_manuale = next((s for s in Stato if s.value == _testo(valore("stato"))), None)
        tempo = _numero(valore("tempogestionesec"))
        giorni = _numero(valore("giorniepisodio"))
        pratiche[protocollo] = Pratica(  # nel file un protocollo è una riga; l'ultima vince
            protocollo=protocollo,
            cf=_testo(valore("cf")).upper(),
            dipendente=_testo(valore("candidato")),
            joborder=_testo(valore("joborder")),
            cliente=_testo(valore("cliente")),
            inizio=_data(valore("datainizio")),
            fine=_data(valore("datafine")),
            data_presentazione=_data(valore("datapresentazione")),
            giorni_episodio=int(giorni) if giorni is not None else None,
            mail_lavoratore=_si(valore("maillav")),
            mail_ufficio=_si(valore("mailuff")),
            stato_manuale=stato_manuale,
            tempo_gestione_sec=tempo,
        )
    return DatiV620("malattie", pratiche=tuple(pratiche.values()), avvisi=tuple(avvisi))


def _giustificativi(colonne: dict[str, int], righe: Any) -> DatiV620:
    lette = []
    avvisi = []
    for numero, riga in enumerate(righe, start=2):
        valore = _lettore(colonne, riga)
        if not any(v not in (None, "") for v in riga):
            continue
        giorno = _data(valore("giorno"))
        cf = _testo(valore("codicefiscale")).upper()
        tipo = _testo(valore("tipogiustificativo"))
        if giorno is None:
            avvisi.append(
                f"riga {numero}: più giorni in una riga, non importata (inseriscila a mano)"
                if _testo(valore("giorno"))
                else f"riga {numero}: manca il giorno, non importata"
            )
            continue
        if not cf or not tipo:
            avvisi.append(f"riga {numero}: mancano codice fiscale o tipo, non importata")
            continue
        ore = _numero(valore("quantitaore"))
        lette.append(
            RigaSalvata(
                cf=cf,
                dipendente=_testo(valore("candidato")),
                joborder=_testo(valore("joborder")),
                cliente=_testo(valore("cliente")),
                giorno=giorno,
                tipo=tipo,
                codice=_testo(valore("codicegiustificativo")),
                minuti=round(ore * 60) if ore is not None else None,
                esportata=True,
            )
        )
    return DatiV620("giustificativi", righe=tuple(lette), avvisi=tuple(avvisi))


def _chiave(intestazione: Any) -> str:
    return normalizza(str(intestazione)).replace(" ", "")


def _lettore(colonne: dict[str, int], riga: Sequence[Any]):
    def valore(nome: str) -> Any:
        posizione = colonne.get(nome)
        return riga[posizione] if posizione is not None and posizione < len(riga) else None

    return valore


def _testo(valore: Any) -> str:
    if valore is None:
        return ""
    if isinstance(valore, float) and valore.is_integer():
        return str(int(valore))
    return str(valore).strip()


def _data(valore: Any) -> date | None:
    if isinstance(valore, datetime):
        return valore.date()
    if isinstance(valore, date):
        return valore
    return leggi_data(_testo(valore)) if valore else None


def _numero(valore: Any) -> float | None:
    testo = _testo(valore).replace(",", ".")
    try:
        return float(testo) if testo else None
    except ValueError:
        return None


def _si(valore: Any) -> bool:
    return normalizza(_testo(valore)) in _SI
