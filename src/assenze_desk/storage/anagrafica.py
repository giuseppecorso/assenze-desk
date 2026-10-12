"""Anagrafica da Excel: dipendenti, orari dei clienti, capienze di ferie e ROL.

Fogli del file:
- "Dipendenti" (obbligatorio): una riga per dipendente e JobOrder. Colonne con
  nomi flessibili (maiuscole, spazi e accenti non contano): CF / Codice fiscale,
  Cognome, Nome (oppure Candidato / Nominativo), Email, JobOrder, Cliente,
  Inizio JobOrder / Data inizio, Fine JobOrder / Data fine.
- "Orari" (facoltativo): Cliente e ore di lavoro da Lunedì a Domenica.
- "Capienze" (facoltativo): JobOrder e ore residue per tipo (Ferie, ROL, ...).
- "Clienti" (facoltativo): Cliente e DL, gli indirizzi del cliente da mettere in
  copia nelle risposte al lavoratore; vince sulla colonna DL dei dipendenti.

Compatibile con il file di configurazione della v6.20 (anche .xlsm): se il
foglio "Dipendenti" non ha le date dei JobOrder, si prendono dal primo altro
foglio con le colonne JobOrder e data di inizio o di fine. Gli altri fogli e
colonne della v6.20 (template, liste di distribuzione) non servono più.

Le righe senza CF si scartano; un CF con il carattere di controllo sbagliato si
importa con un avviso. Negli avvisi il CF è sempre mascherato.
"""

from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, replace
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from assenze_desk.core.codice_fiscale import cf_valido
from assenze_desk.core.date import leggi_data
from assenze_desk.core.dipendenti import Dipendente
from assenze_desk.core.giorni import GIORNI_SETTIMANA
from assenze_desk.core.privacy import maschera_cf
from assenze_desk.core.testo import normalizza
from assenze_desk.core.turni import OrarioSettimanale

FOGLIO_DIPENDENTI = "Dipendenti"
FOGLIO_ORARI = "Orari"
FOGLIO_CAPIENZE = "Capienze"
FOGLIO_CLIENTI = "Clienti"

_ALIAS = {
    "cf": {"cf", "codicefiscale"},
    "cognome": {"cognome"},
    "nome": {"nome"},
    "completo": {"candidato", "dipendente", "nomecompleto", "nominativo"},
    "email": {"email", "mail"},
    "joborder": {"joborder", "numjob", "numjoborder"},
    "cliente": {"cliente", "clienti"},
    "inizio": {
        "iniziojoborder",
        "datainizio",
        "datainiziojo",
        "datainiziojoborder",
        "inizio",
        "dtinizio",
        "datainiziocontratto",
    },
    "copia": {"dl", "distlist", "distributionlist", "copiacliente", "listacliente"},
    "fine": {
        "finejoborder",
        "datafine",
        "datafinejo",
        "datafinejoborder",
        "fine",
        "dtfine",
        "datafinecontratto",
    },
}
_INTESTAZIONI_DIPENDENTI = [
    "CF",
    "Cognome",
    "Nome",
    "Email",
    "JobOrder",
    "Cliente",
    "Inizio JobOrder",
    "Fine JobOrder",
    "DL",
]
_NOMI_GIORNI = ["Lunedì", "Martedì", "Mercoledì", "Giovedì", "Venerdì", "Sabato", "Domenica"]


@dataclass(frozen=True)
class EsitoImport:
    dipendenti: tuple[Dipendente, ...]
    orari: dict[str, OrarioSettimanale]  # cliente → orario
    capienze: dict[str, dict[str, int]]  # JobOrder → tipo → minuti residui
    avvisi: tuple[str, ...]


def importa_anagrafica(percorso: Path) -> EsitoImport:
    if not Path(percorso).is_file():
        raise FileNotFoundError(f"file dell'anagrafica non trovato: {percorso}")
    libro = load_workbook(percorso, read_only=True, data_only=True)
    try:
        if FOGLIO_DIPENDENTI not in libro.sheetnames:
            messaggio = f"nel file manca il foglio '{FOGLIO_DIPENDENTI}'"
            if _sembra_della_v620(libro):
                messaggio += (
                    ": questo è il file delle pratiche o dei giustificativi della versione "
                    "precedente, si importa da Archivio. Qui serve il file con l'elenco "
                    "dei dipendenti"
                )
            raise ValueError(messaggio)
        avvisi: list[str] = []
        dipendenti = _con_periodi(
            _dipendenti(libro[FOGLIO_DIPENDENTI], avvisi), _periodi_da_altri_fogli(libro)
        )
        if FOGLIO_CLIENTI in libro.sheetnames:
            dipendenti = _con_copie(dipendenti, _copie_clienti(libro[FOGLIO_CLIENTI]))
        orari = _orari(libro[FOGLIO_ORARI]) if FOGLIO_ORARI in libro.sheetnames else {}
        capienze = _capienze(libro[FOGLIO_CAPIENZE]) if FOGLIO_CAPIENZE in libro.sheetnames else {}
    finally:
        libro.close()  # in sola lettura openpyxl tiene il file aperto
    return EsitoImport(tuple(dipendenti), orari, capienze, tuple(avvisi))


def scrivi_anagrafica(
    percorso: Path,
    dipendenti: Sequence[Dipendente],
    orari: Mapping[str, OrarioSettimanale],
    capienze: Mapping[str, Mapping[str, int]],
) -> None:
    """Scrive il file nel formato letto da `importa_anagrafica` (anche vuoto, come modello)."""
    libro = Workbook()
    libro.remove(libro.active)

    foglio = _nuovo_foglio(libro, FOGLIO_DIPENDENTI, _INTESTAZIONI_DIPENDENTI)
    for d in dipendenti:
        foglio.append(
            [
                d.cf,
                d.cognome,
                d.nome,
                d.email,
                d.joborder,
                d.cliente,
                d.inizio_joborder,
                d.fine_joborder,
                d.copia_cliente,
            ]
        )
    for cella in (*foglio["G"][1:], *foglio["H"][1:]):
        cella.number_format = "DD/MM/YYYY"

    foglio = _nuovo_foglio(libro, FOGLIO_ORARI, ["Cliente", *_NOMI_GIORNI])
    for cliente, orario in orari.items():
        foglio.append([cliente, *(minuti / 60 for minuti in orario.minuti)])

    tipi = sorted({tipo for per_tipo in capienze.values() for tipo in per_tipo})
    foglio = _nuovo_foglio(libro, FOGLIO_CAPIENZE, ["JobOrder", *tipi])
    for joborder, per_tipo in capienze.items():
        foglio.append([joborder, *(_ore(per_tipo.get(tipo)) for tipo in tipi)])

    Path(percorso).parent.mkdir(parents=True, exist_ok=True)
    libro.save(percorso)


# --- Lettura ---


def _chiave(intestazione: Any) -> str:
    return normalizza(str(intestazione or "")).replace(" ", "")


def _colonne(intestazioni: Sequence[Any]) -> dict[str, int]:
    trovate: dict[str, int] = {}
    for posizione, intestazione in enumerate(intestazioni):
        chiave = _chiave(intestazione)
        for campo, alias in _ALIAS.items():
            if chiave in alias and campo not in trovate:
                trovate[campo] = posizione
    return trovate


def _righe(foglio: Worksheet) -> Iterable[tuple[int, Sequence[Any]]]:
    for numero, riga in enumerate(foglio.iter_rows(min_row=2, values_only=True), start=2):
        if any(valore not in (None, "") for valore in riga):
            yield numero, riga


def _testo(valore: Any) -> str:
    if valore is None:
        return ""
    if isinstance(valore, float) and valore.is_integer():
        return str(int(valore))  # 201111111.0 → "201111111"
    return str(valore).strip()


def _data(valore: Any) -> date | None:
    if isinstance(valore, datetime):  # openpyxl legge le date di Excel come datetime
        return valore.date()
    return leggi_data(_testo(valore)) if valore else None


def _minuti(valore: Any) -> int | None:
    testo = _testo(valore).replace(",", ".")
    try:
        return round(float(testo) * 60) if testo else None
    except ValueError:
        return None


def _ore(minuti: int | None) -> float | None:
    return minuti / 60 if minuti is not None else None


def _dipendenti(foglio: Worksheet, avvisi: list[str]) -> list[Dipendente]:
    intestazioni = next(foglio.iter_rows(max_row=1, values_only=True), ())
    colonne = _colonne(intestazioni)
    if "cf" not in colonne:
        raise ValueError(f"nel foglio '{FOGLIO_DIPENDENTI}' manca la colonna del codice fiscale")

    def campo(riga: Sequence[Any], nome: str) -> Any:
        posizione = colonne.get(nome)
        return riga[posizione] if posizione is not None and posizione < len(riga) else None

    dipendenti = []
    for numero, riga in _righe(foglio):
        cf = _testo(campo(riga, "cf")).upper()
        if not cf:
            avvisi.append(f"riga {numero}: manca il codice fiscale, riga scartata")
            continue
        if not cf_valido(cf):
            avvisi.append(
                f"riga {numero}: codice fiscale {maschera_cf(cf)} non valido "
                "(controlla il carattere finale), riga importata"
            )
        cognome, nome = _testo(campo(riga, "cognome")), _testo(campo(riga, "nome"))
        if not cognome and not nome:
            cognome = _testo(campo(riga, "completo"))
        dipendenti.append(
            Dipendente(
                cf=cf,
                cognome=cognome,
                nome=nome,
                email=_testo(campo(riga, "email")).lower(),
                joborder=_testo(campo(riga, "joborder")),
                cliente=_testo(campo(riga, "cliente")),
                inizio_joborder=_data(campo(riga, "inizio")),
                fine_joborder=_data(campo(riga, "fine")),
                copia_cliente=_testo(campo(riga, "copia")),
            )
        )
    return dipendenti


def _sembra_della_v620(libro: Any) -> bool:
    """Tracking delle malattie o giustificativi della v6.20, scelti per sbaglio."""
    for foglio in libro.worksheets:
        chiavi = {_chiave(h) for h in next(foglio.iter_rows(max_row=1, values_only=True), ())}
        if {"protocollo", "datainizio"} <= chiavi or "tipogiustificativo" in chiavi:
            return True
    return False


def _periodi_da_altri_fogli(libro: Any) -> dict[str, tuple[date | None, date | None]]:
    """JobOrder → (inizio, fine) dal primo foglio che li contiene, come nella v6.20."""
    for nome in libro.sheetnames:
        if nome in (FOGLIO_DIPENDENTI, FOGLIO_ORARI, FOGLIO_CAPIENZE, FOGLIO_CLIENTI):
            continue
        foglio = libro[nome]
        colonne = _colonne(next(foglio.iter_rows(max_row=1, values_only=True), ()))
        if "joborder" not in colonne or not {"inizio", "fine"} & colonne.keys():
            continue

        def campo(riga: Sequence[Any], chiave: str, colonne=colonne) -> Any:
            posizione = colonne.get(chiave)
            return riga[posizione] if posizione is not None and posizione < len(riga) else None

        periodi = {}
        for _numero, riga in _righe(foglio):
            joborder = _testo(campo(riga, "joborder"))
            inizio, fine = _data(campo(riga, "inizio")), _data(campo(riga, "fine"))
            if joborder and (inizio or fine):
                periodi[joborder] = (inizio, fine)
        if periodi:
            return periodi
    return {}


def _con_periodi(
    dipendenti: list[Dipendente], periodi: Mapping[str, tuple[date | None, date | None]]
) -> list[Dipendente]:
    """Le date scritte nel foglio Dipendenti vincono; le mancanti vengono dai periodi."""
    completati = []
    for d in dipendenti:
        inizio, fine = periodi.get(d.joborder, (None, None))
        completati.append(
            replace(
                d,
                inizio_joborder=d.inizio_joborder or inizio,
                fine_joborder=d.fine_joborder or fine,
            )
        )
    return completati


def _copie_clienti(foglio: Worksheet) -> dict[str, str]:
    """Cliente (normalizzato) → indirizzi da mettere in copia."""
    colonne = _colonne(next(foglio.iter_rows(max_row=1, values_only=True), ()))
    if not {"cliente", "copia"} <= colonne.keys():
        return {}
    copie = {}
    for _numero, riga in _righe(foglio):
        cliente = normalizza(_testo(riga[colonne["cliente"]]))
        copia = _testo(riga[colonne["copia"]]) if colonne["copia"] < len(riga) else ""
        if cliente and copia:
            copie[cliente] = copia
    return copie


def _con_copie(dipendenti: list[Dipendente], copie: Mapping[str, str]) -> list[Dipendente]:
    return [
        replace(d, copia_cliente=copie.get(normalizza(d.cliente), d.copia_cliente))
        for d in dipendenti
    ]


def _orari(foglio: Worksheet) -> dict[str, OrarioSettimanale]:
    intestazioni = [_chiave(h) for h in next(foglio.iter_rows(max_row=1, values_only=True), ())]
    if "cliente" not in intestazioni:
        return {}
    posizione_cliente = intestazioni.index("cliente")
    posizioni_giorni = [
        intestazioni.index(giorno) if giorno in intestazioni else None
        for giorno in GIORNI_SETTIMANA
    ]
    orari = {}
    for _numero, riga in _righe(foglio):
        cliente = _testo(riga[posizione_cliente])
        minuti = tuple(
            (_minuti(riga[p]) or 0) if p is not None and p < len(riga) else 0
            for p in posizioni_giorni
        )
        if cliente:
            orari[cliente] = OrarioSettimanale(minuti)
    return orari


def _capienze(foglio: Worksheet) -> dict[str, dict[str, int]]:
    intestazioni = list(next(foglio.iter_rows(max_row=1, values_only=True), ()))
    chiavi = [_chiave(h) for h in intestazioni]
    if "joborder" not in chiavi:
        return {}
    posizione_job = chiavi.index("joborder")
    capienze: dict[str, dict[str, int]] = {}
    for _numero, riga in _righe(foglio):
        joborder = _testo(riga[posizione_job])
        per_tipo = {
            _testo(intestazioni[p]): minuti
            for p in range(len(riga))
            if p != posizione_job and (minuti := _minuti(riga[p])) is not None
        }
        if joborder and per_tipo:
            capienze[joborder] = per_tipo
    return capienze


# --- Scrittura ---


def _nuovo_foglio(libro: Workbook, nome: str, intestazioni: list[str]) -> Worksheet:
    foglio = libro.create_sheet(nome)
    foglio.append(intestazioni)
    for cella in foglio[1]:
        cella.font = Font(bold=True)
    foglio.freeze_panes = "A2"
    for colonna in range(1, len(intestazioni) + 1):
        foglio.column_dimensions[foglio.cell(1, colonna).column_letter].width = 18
    return foglio
